"""Reopen generated files and check content and Excel validation before commit."""
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from jinsey.errors import JinseyError
from jinsey.exporters.excel_exporter import COLS, cell_value
from jinsey.schema import CaseSheet, ChoiceField, TestCase


def verify_yaml(path: Path, cases: list[TestCase]) -> None:
    with path.open(encoding="utf-8") as handle:
        actual = yaml.safe_load(handle)
    if actual != {"test_cases": [case.to_dict() for case in cases]}:
        raise JinseyError(f"YAML verification failed: {path}")


def verify_yaml_sheets(path: Path, sheets: list[CaseSheet]) -> None:
    with path.open(encoding="utf-8") as handle:
        actual = yaml.safe_load(handle)
    if actual != {"sheets": [sheet.to_dict() for sheet in sheets]}:
        raise JinseyError(f"YAML verification failed: {path}")


def _verify_sheet(wb: Workbook, ws: Worksheet, cases: list[TestCase], fields: dict[str, ChoiceField]) -> None:
    expected_headers = [header for _, header, _ in COLS]
    if [c.value for c in ws[1]] != expected_headers or ws.max_row != len(cases) + 1:
        raise JinseyError("Excel verification failed: headers or row count")
    for row, case in zip(ws.iter_rows(min_row=2), cases, strict=True):
        for cell, (field, _, _) in zip(row, COLS, strict=True):
            expected = cell_value(getattr(case, field))
            if (cell.value or "") != expected or cell.data_type == "f" or not cell.alignment.wrap_text:
                raise JinseyError(f"Excel verification failed: {case.key}, {field}")
    column = ws.cell(1, expected_headers.index("Priority") + 1).column_letter
    validations = list(ws.data_validations.dataValidation)
    if len(validations) != 1:
        raise JinseyError("Excel verification failed: missing Priority validation")
    dv = validations[0]
    if (dv.type != "list" or not dv.showErrorMessage or dv.errorStyle != "stop"
            or dv.allow_blank or dv.showDropDown
            or any(f"{column}{row}" not in dv for row in range(2, len(cases) + 2))):
        raise JinseyError("Excel verification failed: Priority dropdown configuration")
    if dv.formula1 == "PriorityOptions":
        choices = tuple(row[0].value for row in wb["_choices"].iter_rows())
    else:
        choices = tuple(dv.formula1.strip('"').split(","))
    if choices != fields["priority"].options or ws.freeze_panes != "A2" or ws.auto_filter.ref != ws.dimensions:
        raise JinseyError("Excel verification failed: choices, frozen header, or filters")


def verify_excel_sheets(path: Path, sheets: list[CaseSheet], fields: dict[str, ChoiceField]) -> None:
    wb = load_workbook(path)
    try:
        visible = [ws.title for ws in wb if ws.sheet_state == "visible"]
        if visible != [sheet.name for sheet in sheets]:
            raise JinseyError("Excel verification failed: sheet names/order")
        if set(wb.sheetnames) - {sheet.name for sheet in sheets} - {"_choices"}:
            raise JinseyError("Excel verification failed: unexpected sheets")
        for sheet in sheets:
            _verify_sheet(wb, wb[sheet.name], sheet.cases, fields)
    finally:
        wb.close()


def verify_excel(path: Path, cases: list[TestCase], fields: dict[str, ChoiceField]) -> None:
    verify_excel_sheets(path, [CaseSheet("Test Cases", "", cases)], fields)
