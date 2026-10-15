import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from jinsey.schema import CaseSheet, ChoiceField, TestCase, validate_sheet_names

COLS = [("id", "ID", 16), ("key", "Key", 32), ("feature", "Feature", 16),
        ("scenario", "Scenario", 48), ("precondition", "Preconditions", 45),
        ("steps", "Steps", 60), ("expected", "Expected Result", 55),
        ("priority", "Priority", 14), ("tags", "Tags", 25),
        ("test_data", "Test Data", 40), ("test_type", "Test Type", 20)]

INVALID_XML_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
MAX_DATA_ROWS = 1_048_575


def cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, tuple):
        return ", ".join(value)
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _populate_sheet(wb: Workbook, ws: Worksheet, cases: list[TestCase], fields: dict[str, ChoiceField]) -> None:
    if len(cases) > MAX_DATA_ROWS:
        raise ValueError(f"Excel supports at most {MAX_DATA_ROWS} test cases plus the header")
    ws.append([header for _, header, _ in COLS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="263D57")
    wrapped = Alignment(wrap_text=True, vertical="top")
    for row_index, case in enumerate(cases, 2):
        values = [cell_value(getattr(case, field)) for field, _, _ in COLS]
        # Excel's cell limit must fail explicitly instead of silently truncating content.
        if any(len(value) > 32767 for value in values):
            raise ValueError(f"Test case {case.key}: Excel cells cannot exceed 32767 characters")
        for value, (field, _, _) in zip(values, COLS, strict=True):
            if INVALID_XML_TEXT.search(value):
                raise ValueError(f"Test case {case.key}: {field} contains characters forbidden in Excel XML")
        ws.append(values)
        for column in range(1, len(COLS) + 1):
            cell = ws.cell(row_index, column)
            cell.data_type = "s"  # Preserve IDs and prevent authored formulas from executing.
            cell.alignment = wrapped
        ws.row_dimensions[row_index].height = min(409, max(45, 15 * max(v.count("\n") + 1 for v in values)))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, (_, _, width) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    headers = {cell.value: cell.column_letter for cell in ws[1]}
    options = fields["priority"].options
    inline = ",".join(options)
    if len(inline) <= 253 and all(not any(c in option for c in ',"\n\r') for option in options):
        formula = f'"{inline}"'
    elif "_choices" not in wb.sheetnames:
        lookup = wb.create_sheet("_choices")
        for index, option in enumerate(options, 1):
            lookup.append([option])
            lookup.cell(index, 1).data_type = "s"
        lookup.sheet_state = "hidden"
        wb.defined_names.add(DefinedName("PriorityOptions", attr_text=f"'_choices'!$A$1:$A${len(options)}"))
        formula = "PriorityOptions"
    else:
        formula = "PriorityOptions"
    validation = DataValidation(type="list", formula1=formula, allow_blank=False)
    validation.showDropDown = False
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    validation.errorTitle = "Invalid Priority"
    validation.error = "Choose one of the configured Priority options."
    validation.showInputMessage = True
    validation.promptTitle = "Priority"
    validation.prompt = "Select a priority from the dropdown."
    ws.add_data_validation(validation)
    validation.add(f"{headers['Priority']}2:{headers['Priority']}{max(2, len(cases) + 1)}")


def export_excel_sheets(sheets: list[CaseSheet], path: Path, fields: dict[str, ChoiceField]) -> None:
    validate_sheet_names([sheet.name for sheet in sheets])
    wb = Workbook()
    wb.remove(wb.active)
    try:
        for sheet in sheets:
            ws = wb.create_sheet(sheet.name)
            _populate_sheet(wb, ws, sheet.cases, fields)
        wb.save(path)
    finally:
        wb.close()


def export_excel(cases: list[TestCase], path: Path, fields: dict[str, ChoiceField]) -> None:
    export_excel_sheets([CaseSheet("Test Cases", "", cases)], path, fields)
