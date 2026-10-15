import json

import pytest
import yaml
from openpyxl import load_workbook

from jinsey.errors import JinseyError
from jinsey.generator import generate, validate
from jinsey.schema import load_fields
from jinsey.verification import verify_excel, verify_yaml


def test_end_to_end_ten_cases(root, tmp_path):
    template = root / "data/input/testcases.j2"
    config = root / "config/fields.yaml"
    registry = tmp_path / "registry.json"
    output = tmp_path / "output"
    result = generate(template, config, output, registry, prefix="TC")
    assert len(result.cases) == 10
    assert [case.id for case in result.cases] == [f"TC-{i:03d}" for i in range(1, 11)]
    assert len({len(case.steps.splitlines()) for case in result.cases}) >= 3
    verify_yaml(output / "testcases.yaml", result.cases)
    verify_excel(output / "testcases.xlsx", result.cases, load_fields(config))
    before = registry.read_bytes()
    second = generate(template, config, output, registry, prefix="TC")
    assert second.cases == result.cases and registry.read_bytes() == before
    wb = load_workbook(output / "testcases.xlsx")
    assert wb.active.max_row == 11
    wb.close()
    auth = generate(template, config, output, registry, prefix="AUTH")
    assert auth.cases[0].id == "AUTH-001"


def test_editing_template_updates_both_formats(root, tmp_path, template, raw_case):
    config = root / "config/fields.yaml"
    output, registry = tmp_path / "output", tmp_path / "registry.json"
    first = generate(template, config, output, registry)
    raw_case["steps"] = "Langkah baru\nSelesaikan perubahan\nPeriksa hasil"
    template.write_text(json.dumps({"test_cases": [raw_case]}))
    updated = generate(template, config, output, registry)
    assert updated.cases[0].id == first.cases[0].id
    assert "1. Langkah baru" in yaml.safe_load((output / "testcases.yaml").read_text())["test_cases"][0]["steps"]
    verify_excel(output / "testcases.xlsx", updated.cases, load_fields(config))


@pytest.mark.parametrize("format,filename", [("yaml", "testcases.yaml"), ("excel", "testcases.xlsx")])
def test_single_format(root, tmp_path, template, format, filename):
    result = generate(template, root / "config/fields.yaml", tmp_path / "output", tmp_path / "registry.json", format)
    assert [p.name for p in result.artifacts] == [filename]
    assert [p.name for p in (tmp_path / "output").iterdir()] == [filename]


def test_validate_and_dry_run_are_readonly(root, tmp_path, template):
    config = root / "config/fields.yaml"
    before = set(tmp_path.rglob("*"))
    assert len(validate(template, config)[1]) == 1
    result = generate(template, config, tmp_path / "out", tmp_path / "state" / "ids.json", dry_run=True)
    assert result.cases[0].id == "TC-001"
    assert set(tmp_path.rglob("*")) == before


@pytest.mark.parametrize("failure", ["export", "verification", "registry"])
def test_failures_preserve_outputs_and_registry(root, tmp_path, template, monkeypatch, failure):
    config, output, registry = root / "config/fields.yaml", tmp_path / "output", tmp_path / "registry.json"
    result = generate(template, config, output, registry)
    before = {p: p.read_bytes() for p in [*result.artifacts, registry]}

    def fail(*args):
        raise OSError("simulated failure")

    target = {"export": "jinsey.generator.export_excel", "verification": "jinsey.generator.verify_excel",
              "registry": "jinsey.id_registry.IDRegistry.write"}[failure]
    monkeypatch.setattr(target, fail)
    with pytest.raises(OSError, match="simulated"):
        generate(template, config, output, registry)
    assert all(p.read_bytes() == data for p, data in before.items())


def test_invalid_input_no_side_effect(root, tmp_path, template):
    template.write_text('{"test_cases": []}')
    with pytest.raises(JinseyError, match="nonempty"):
        generate(template, root / "config/fields.yaml", tmp_path / "out", tmp_path / "registry.json")
    assert not (tmp_path / "out").exists() and not (tmp_path / "registry.json").exists()


def test_invalid_format(root, tmp_path, template):
    with pytest.raises(JinseyError, match="Unsupported format"):
        generate(template, root / "config/fields.yaml", tmp_path / "out", tmp_path / "registry.json", format="csv")


def test_protect_authoring_files(root, tmp_path, template):
    before = template.read_bytes()
    with pytest.raises(JinseyError, match="overwrite"):
        generate(template, root / "config/fields.yaml", tmp_path / "out", template)
    assert template.read_bytes() == before
