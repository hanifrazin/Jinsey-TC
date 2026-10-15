import json

import pytest
from openpyxl import load_workbook

from jinsey.errors import JinseyError
from jinsey.generator import generate
from jinsey.processor import normalize


def test_dynamic_dropdown_generation(tmp_path, raw_case, root):
    template = tmp_path / "cases.j2"
    raw_cases = [dict(raw_case, key=f"case-{i}", dropdown=dropdown) for i, dropdown in enumerate([
        {"Kategori": "A", "Status": {"value": "Ready", "options": ["Ready", "Later, maybe"]}},
        {"Kategori": "B", "Label": "=1+1"},
        {"Kategori": "C"},
        {},
    ])]
    template.write_text(json.dumps({"test_cases": raw_cases}))
    result = generate(template, root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json")
    assert result.cases[0].to_dict()["dropdown"]["Kategori"] == "A"
    wb = load_workbook(tmp_path / "out" / "testcases.xlsx")
    ws = wb.active
    columns = {cell.value: cell.column for cell in ws[1]}
    assert [ws.cell(i, columns["Kategori"]).value for i in range(2, 6)] == ["A", "B", "C", None]
    dv = list(ws.data_validations.dataValidation)
    assert dv[1].formula1 == '"A,B,C"'
    assert dv[1].allow_blank
    assert str(dv[1].sqref) == "L2:L5"
    assert wb["_choices"].sheet_state == "hidden"
    assert ws.cell(3, columns["Label"]).data_type == "s"
    wb.close()


@pytest.mark.parametrize("dropdown", [[], {"Priority": "A"}, {"Kategori": ""},
    {"Kategori": "A", "kategori": "B"}, {" Kategori": "A"},
    {"Kategori": {"value": "D", "options": ["A", "B"]}},
    {"Kategori": {"value": "A", "options": ["A", "A"]}},
    {"Kategori": {"value": "A", "options": [1]}},
    {"Kategori": {"value": "A", "options": ["A"], "extra": True}},
    {"Kategori": "bad\x00text"}, {"Kategori": "x" * 32768},
])
def test_invalid_dropdown(raw_case, fields, dropdown):
    with pytest.raises(JinseyError, match="dropdown"):
        normalize({"test_cases": [dict(raw_case, dropdown=dropdown)]}, fields)


def test_dropdowns_are_scoped_per_sheet(tmp_path, raw_case, root):
    suite = tmp_path / "suite"
    suite.mkdir()
    for name, options in [("one", ["A,1", "B"]), ("two", ["C", 'D"2'])]:
        case = dict(raw_case, key=name, dropdown={"Kategori": {"value": options[0], "options": options}})
        (suite / f"{name}.j2").write_text(json.dumps({"test_cases": [case]}))
    generate(suite, root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json")
    wb = load_workbook(tmp_path / "out" / "testcases.xlsx")
    assert wb["one"].data_validations.dataValidation[1].formula1 != wb["two"].data_validations.dataValidation[1].formula1
    wb.close()
