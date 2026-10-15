"""Folder input produces one transaction with one Excel sheet per template."""
import json
from dataclasses import replace
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from jinsey.cli import main
from jinsey.errors import JinseyError
from jinsey.exporters.excel_exporter import export_excel_sheets
from jinsey.generator import generate, validate
from jinsey.schema import ChoiceField, load_fields, validate_sheet_names
from jinsey.verification import verify_excel_sheets, verify_yaml_sheets


@pytest.fixture
def suite(tmp_path, raw_case):
    folder = tmp_path / "suite"
    folder.mkdir()
    # Intentionally create in reverse order to verify deterministic discovery.
    for index in reversed(range(1, 6)):
        cases = [{**raw_case, "key": f"case-{index}-{n}", "prefix": "LOGIN"} for n in range(index)]
        (folder / f"part_{index}.j2").write_text(json.dumps({"test_cases": cases}))
    (folder / "README.txt").write_text("Not a template")
    shared = folder / "shared"
    shared.mkdir()
    (shared / "helper.j2").write_text("This helper is not a test case envelope")
    (folder / ".editor.j2").write_text("Ignore hidden editor files")
    return folder


def test_five_files_become_five_sheets(root, tmp_path, suite):
    config = root / "config/fields.yaml"
    output, registry = tmp_path / "out", tmp_path / "ids.json"
    result = generate(suite, config, output, registry)
    assert len(result.cases) == 15
    assert [sheet.name for sheet in result.sheets] == [f"part_{i}" for i in range(1, 6)]
    assert [len(sheet.cases) for sheet in result.sheets] == [1, 2, 3, 4, 5]
    assert [case.id for case in result.cases] == [f"LOGIN-{i:03d}" for i in range(1, 16)]
    wb = load_workbook(output / "testcases.xlsx")
    assert wb.sheetnames == [sheet.name for sheet in result.sheets]
    for sheet in result.sheets:
        ws = wb[sheet.name]
        assert ws.max_row == len(sheet.cases) + 1
        assert ws["A2"].value == sheet.cases[0].id
        assert ws["F2"].value == sheet.cases[0].steps
        assert ws["F2"].alignment.wrap_text
        assert ws.freeze_panes == "A2"
        dv = list(ws.data_validations.dataValidation)[0]
        assert dv.formula1 == '"P0,P1,P2"' and dv.showErrorMessage
    wb.close()
    data = yaml.safe_load((output / "testcases.yaml").read_text())
    assert [item["name"] for item in data["sheets"]] == [sheet.name for sheet in result.sheets]
    assert [item["source"] for item in data["sheets"]] == [f"part_{i}.j2" for i in range(1, 6)]
    verify_yaml_sheets(output / "testcases.yaml", result.sheets)
    verify_excel_sheets(output / "testcases.xlsx", result.sheets, load_fields(config))
    before = registry.read_bytes()
    assert generate(suite, config, output, registry).cases == result.cases
    assert registry.read_bytes() == before


def test_folder_validate_and_dry_run_do_not_write(root, tmp_path, suite):
    before = set(tmp_path.rglob("*"))
    config = root / "config/fields.yaml"
    assert len(validate(suite, config)[1]) == 15
    preview = generate(suite, config, tmp_path / "out", tmp_path / "state/ids.json", dry_run=True)
    assert len(preview.sheets) == 5
    assert set(tmp_path.rglob("*")) == before


def test_duplicate_across_files_fails_before_writing(root, tmp_path, suite):
    first = suite / "part_1.j2"
    (suite / "part_5.j2").write_bytes(first.read_bytes())
    before = set(tmp_path.rglob("*"))
    with pytest.raises(JinseyError, match="Duplicate key.*part_1.j2.*part_5.j2"):
        generate(suite, root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json")
    assert set(tmp_path.rglob("*")) == before


@pytest.mark.parametrize("source", ['{% if %}', '{"test_cases": []}'])
def test_invalid_file_names_source_and_preserves_previous_outputs(root, tmp_path, suite, source):
    config, output, registry = root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json"
    result = generate(suite, config, output, registry)
    before = {path: path.read_bytes() for path in [*result.artifacts, registry]}
    (suite / "part_5.j2").write_text(source)
    with pytest.raises(JinseyError, match="part_5.j2"):
        generate(suite, config, output, registry)
    assert all(path.read_bytes() == contents for path, contents in before.items())


@pytest.mark.parametrize("names", [["a" * 32], ["bad:name"], ["bad/name"], ["bad\\name"],
    ["[name]"], ["bad?name"], ["bad*name"], ["'name"], ["name'"], [""], ["\x00"],
    ["Login", "login"], ["_choices"], []])
def test_invalid_sheet_names(names):
    with pytest.raises(JinseyError, match="sheet|Sheet|Workbook"):
        validate_sheet_names(names)


def test_long_filename_rejected_before_writing(root, tmp_path, suite):
    (suite / "part_1.j2").rename(suite / ("a" * 32 + ".j2"))
    with pytest.raises(JinseyError, match="Invalid sheet name"):
        generate(suite, root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json")
    assert not (tmp_path / "out").exists()


def test_empty_folder_reports_helpful_error(root, tmp_path):
    with pytest.raises(JinseyError, match="No .j2 files directly inside"):
        validate(tmp_path, root / "config/fields.yaml")


@pytest.mark.parametrize("format,filename", [("yaml", "testcases.yaml"), ("excel", "testcases.xlsx")])
def test_folder_single_format(root, tmp_path, suite, format, filename):
    result = generate(suite, root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json", format)
    assert [path.name for path in result.artifacts] == [filename]
    assert [path.name for path in (tmp_path / "out").iterdir()] == [filename]


def test_folder_override_prefix(root, tmp_path, suite):
    result = generate(suite, root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json", prefix="AUTH")
    assert [case.id for case in result.cases] == [f"AUTH-{i:03d}" for i in range(1, 16)]


def test_folder_rename_insert_and_move_preserve_ids(root, tmp_path, suite, raw_case):
    config, output, registry = root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json"
    original = generate(suite, config, output, registry)
    known = {case.key: case.id for case in original.cases}
    (suite / "part_1.j2").rename(suite / "z_last.j2")
    (suite / "a_first.j2").write_text(json.dumps({"test_cases": [{**raw_case, "key": "new-first", "prefix": "LOGIN"}]}))
    first = json.loads((suite / "part_2.j2").read_text())
    second = json.loads((suite / "part_3.j2").read_text())
    second["test_cases"].append(first["test_cases"].pop())
    (suite / "part_2.j2").write_text(json.dumps(first))
    (suite / "part_3.j2").write_text(json.dumps(second))
    updated = generate(suite, config, output, registry)
    assert updated.sheets[0].name == "a_first" and updated.sheets[-1].name == "z_last"
    assert updated.cases[0].id == "LOGIN-016"
    assert {case.key: case.id for case in updated.cases if case.key in known} == known


def test_each_file_has_its_own_prefix_fallback(root, tmp_path, raw_case):
    folder = tmp_path / "suite"
    folder.mkdir()
    (folder / "one.j2").write_text(json.dumps({"prefix": "AUTH", "test_cases": [raw_case]}))
    (folder / "two.j2").write_text(json.dumps({"test_cases": [{**raw_case, "key": "second"}]}))
    result = generate(folder, root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json")
    assert [case.id for case in result.cases] == ["AUTH-001", "TC-001"]


def test_helpers_inside_suite_are_imported_without_becoming_sheets(root, tmp_path, suite, raw_case):
    (suite / "shared/helper.j2").write_text('{% macro active() %}Fictional active account{% endmacro %}')
    case = {**raw_case, "key": "uses-shared"}
    source = '{% import "shared/helper.j2" as pre %}\n{% set case = ' + json.dumps(case) + ' %}\n'
    source += '{% set _ = case.update({"precondition": pre.active()}) %}\n{{ {"test_cases": [case]} | tojson }}'
    (suite / "part_1.j2").write_text(source)
    result = generate(suite, root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json")
    assert len(result.sheets) == 5
    assert result.sheets[0].cases[0].precondition == "Fictional active account"


def test_shared_hidden_dropdown_sheet(root, tmp_path, suite):
    config = root / "config/fields.yaml"
    preview = generate(suite, config, tmp_path / "out", tmp_path / "ids.json", dry_run=True)
    options = ("a" * 130, "b" * 130)
    fields = {"priority": ChoiceField(options, options[0])}
    sheets = [replace(sheet, cases=[replace(case, priority=options[0]) for case in sheet.cases])
              for sheet in preview.sheets]
    path = tmp_path / "choices.xlsx"
    export_excel_sheets(sheets, path, fields)
    verify_excel_sheets(path, sheets, fields)
    wb = load_workbook(path)
    assert len(wb.sheetnames) == 6
    assert wb["_choices"].sheet_state == "hidden"
    assert list(wb.defined_names) == ["PriorityOptions"]
    for sheet in sheets:
        assert list(wb[sheet.name].data_validations.dataValidation)[0].formula1 == "PriorityOptions"
    wb.close()


@pytest.mark.parametrize("target", ["export_excel_sheets", "verify_excel_sheets", "verify_yaml_sheets"])
def test_folder_failure_rolls_back(root, tmp_path, suite, monkeypatch, target):
    config, output, registry = root / "config/fields.yaml", tmp_path / "out", tmp_path / "ids.json"
    result = generate(suite, config, output, registry)
    before = {path: path.read_bytes() for path in [*result.artifacts, registry]}

    def fail(*args):
        raise OSError("simulated folder export failure")

    monkeypatch.setattr(f"jinsey.generator.{target}", fail)
    with pytest.raises(OSError, match="simulated"):
        generate(suite, config, output, registry)
    assert all(path.read_bytes() == contents for path, contents in before.items())


@pytest.mark.parametrize("flag,value", [("--input-dir", "auth_suite"), ("--template", "auth_suite"),
    ("--input-dir", "data/input/auth_suite"), ("--input-dir", "nested/auth_suite")])
def test_cli_folder_paths(root, tmp_path, suite, monkeypatch, capsys, flag, value):
    destination = tmp_path / "data/input" / ("nested/auth_suite" if value.startswith("nested") else "auth_suite")
    destination.parent.mkdir(parents=True)
    suite.rename(destination)
    monkeypatch.chdir(tmp_path)
    assert main(["generate", flag, value, "--fields", str(root / "config/fields.yaml")]) == 0
    stdout = capsys.readouterr().out
    assert "Input directory:" in stdout and "Sheets: part_1 (1 TC)" in stdout
    assert Path("data/output/auth_suite/testcases.xlsx").exists()


def test_cli_input_dir_rejects_file(root, tmp_path, suite, capsys):
    assert main(["generate", "--input-dir", str(suite / "part_1.j2"),
                 "--fields", str(root / "config/fields.yaml")]) == 1
    assert "must refer to a folder" in capsys.readouterr().err


def test_cli_inputs_mutually_exclusive():
    with pytest.raises(SystemExit) as exc:
        main(["generate", "--input-dir", "folder", "--template", "cases.j2"])
    assert exc.value.code == 2


def test_folder_symlink_escape_rejected(root, tmp_path, suite):
    external = tmp_path / "external.j2"
    external.write_bytes((suite / "part_1.j2").read_bytes())
    try:
        (suite / "leak.j2").symlink_to(external)
    except OSError:
        pytest.skip("Symlink creation unavailable")
    with pytest.raises(JinseyError, match="outside input directory"):
        validate(suite, root / "config/fields.yaml")


def test_five_example_files_cover_same_cases(root, tmp_path):
    config = root / "config/fields.yaml"
    result = generate(root / "data/input/auth_suite", config, tmp_path / "out", tmp_path / "ids.json")
    assert len(result.sheets) == 5 and len(result.cases) == 10
    original = validate(root / "data/input/testcases.j2", config)[1]
    assert {case.key: replace(case, id="") for case in result.cases} == {case.key: case for case in original}
