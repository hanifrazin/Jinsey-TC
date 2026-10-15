from dataclasses import replace

import pytest
import yaml
from openpyxl import load_workbook

from jinsey.errors import JinseyError
from jinsey.exporters.excel_exporter import export_excel
from jinsey.exporters.yaml_exporter import export_yaml
from jinsey.processor import normalize
from jinsey.schema import ChoiceField
from jinsey.verification import verify_excel, verify_yaml


def test_roundtrip_multiline_strings_and_dropdown(tmp_path, raw_case, fields):
    case = normalize({"test_cases": [raw_case]}, fields)[1][0].with_id("TC-001")
    yml, xlsx = tmp_path / "cases.yaml", tmp_path / "cases.xlsx"
    export_yaml([case], yml)
    export_excel([case], xlsx, fields)
    verify_yaml(yml, [case])
    verify_excel(xlsx, [case], fields)
    assert "steps: |-" in yml.read_text()
    assert yaml.safe_load(yml.read_text())["test_cases"][0]["id"] == "TC-001"
    wb = load_workbook(xlsx)
    ws = wb.active
    columns = {cell.value: cell.column for cell in ws[1]}
    assert ws.cell(2, columns["ID"]).data_type == "s"
    assert ws.cell(2, columns["Steps"]).value == case.steps
    assert ws.cell(2, columns["Steps"]).alignment.wrap_text
    dv = list(ws.data_validations.dataValidation)[0]
    assert dv.formula1 == '"P0,P1,P2"'
    assert dv.showErrorMessage and dv.error and dv.showDropDown is False
    assert str(dv.sqref) == f"{ws.cell(1, columns['Priority']).column_letter}2"
    wb.close()


@pytest.mark.parametrize("options", [("a" * 130, "b" * 130), ("P,0", "P1"), ('P"0', "P1")])
def test_long_and_special_options_use_lookup(tmp_path, raw_case, fields, options):
    case = normalize({"test_cases": [raw_case]}, fields)[1][0].with_id("TC-001")
    fields = {"priority": ChoiceField(options, options[0])}
    path = tmp_path / "lookup.xlsx"
    export_excel([replace(case, priority=options[0])], path, fields)
    verify_excel(path, [replace(case, priority=options[0])], fields)
    wb = load_workbook(path)
    assert wb["_choices"].sheet_state == "hidden"
    assert "PriorityOptions" in wb.defined_names
    wb.close()


def test_formula_text_and_structured_data(tmp_path, raw_case, fields):
    case = normalize({"test_cases": [raw_case]}, fields)[1][0]
    case = replace(case, scenario="=1+1", tags=("smoke", "fictional"), test_data={"n": [1, True]})
    path = tmp_path / "formula.xlsx"
    export_excel([case], path, fields)
    verify_excel(path, [case], fields)
    wb = load_workbook(path)
    assert wb.active["D2"].value == "=1+1" and wb.active["D2"].data_type == "s"
    wb.close()


def test_excel_limits_fail(tmp_path, raw_case, fields):
    case = normalize({"test_cases": [raw_case]}, fields)[1][0]
    with pytest.raises(ValueError, match="32767"):
        export_excel([replace(case, expected="x" * 32768)], tmp_path / "big.xlsx", fields)


def test_verification_detects_corruption(tmp_path, raw_case, fields):
    cases = normalize({"test_cases": [raw_case]}, fields)[1]
    path = tmp_path / "bad.yaml"
    path.write_text("test_cases: []")
    with pytest.raises(JinseyError, match="YAML verification"):
        verify_yaml(path, cases)
    path = tmp_path / "bad.xlsx"
    export_excel(cases, path, fields)
    wb = load_workbook(path)
    wb.active["F2"] = "Changed"
    wb.save(path)
    wb.close()
    with pytest.raises(JinseyError, match="Excel verification"):
        verify_excel(path, cases, fields)
